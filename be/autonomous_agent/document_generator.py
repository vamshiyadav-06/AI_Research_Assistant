import os
import re
from datetime import datetime
from typing import List

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from .schemas import ExecutorOutput, PlannerOutput

GENERATED_DOCS_DIR = "generated_docs"


def _slugify(value: str) -> str:
    slug = re.sub(r"[^a-zA-Z0-9]+", "_", value.strip().lower()).strip("_")
    return slug or "document"


def _add_bullets(document: Document, items: List[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def generate_word_document(plan: PlannerOutput, content: ExecutorOutput) -> str:
    os.makedirs(GENERATED_DOCS_DIR, exist_ok=True)

    document = Document()
    styles = document.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)

    title = document.add_heading(content.title, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_paragraph(f"Document Type: {plan.document_type}")
    document.add_paragraph(f"Generated On: {datetime.now().strftime('%Y-%m-%d %H:%M')}")

    document.add_heading("Assumptions", level=1)
    _add_bullets(document, plan.assumptions)

    for section in content.sections:
        document.add_heading(section.heading, level=1)
        for paragraph in section.content.split("\n"):
            paragraph = paragraph.strip()

            if paragraph:
                document.add_paragraph(paragraph)

    filename = f"{_slugify(plan.document_type)}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
    path = os.path.join(GENERATED_DOCS_DIR, filename)
    document.save(path)

    return path.replace("\\", "/")
